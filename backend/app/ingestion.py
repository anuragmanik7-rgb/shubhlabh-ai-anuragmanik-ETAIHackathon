from io import BytesIO, StringIO
from datetime import datetime, timezone, timedelta
from decimal import Decimal
import csv, json, re, zipfile
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image
from .domain import AppError, uid, digest, now, CITIES,business_date
from .providers import Provider
from .economics import rounded
from .execution import check_inventory_observation

def table_rows(name,content):
    book=None
    try:
        if name.lower().endswith('.csv'):
            reader=iter(csv.reader(StringIO(content.decode('utf-8-sig')),strict=True))
        elif name.lower().endswith('.xlsx'):
            with zipfile.ZipFile(BytesIO(content)) as z:
                if sum(f.file_size for f in z.infolist())>40_000_000:raise AppError('EXPANDED_LIMIT','Spreadsheet expansion is too large')
            book=load_workbook(BytesIO(content),read_only=True,data_only=True)
            if (book.active.max_row is not None and book.active.max_row>10001) or (book.active.max_column is not None and book.active.max_column>60):raise AppError('ROW_LIMIT','Use up to 10,000 rows and 60 columns')
            # Dimensions are optional and can understate the actual XML. Read
            # the real stream, with an explicit physical-row bound below.
            book.active.reset_dimensions()
            reader=iter(book.active.iter_rows(values_only=True))
        else:raise AppError('FILE_TYPE','Use CSV or XLSX for operational tables')
        headers=next(reader,None)
        if not headers or any(not isinstance(h,str) or not h.strip() for h in headers):raise AppError('TABLE_HEADER','A nonempty header is required for every column')
        headers=[h.strip() for h in headers]
        if len(headers)>60 or len(headers)!=len(set(headers)):raise AppError('TABLE_HEADER','Column names must be unique; use at most 60 columns')
        out=[]
        for row_number,row in enumerate(reader,2):
            if row_number>10001:raise AppError('ROW_LIMIT','Use up to 10,000 data rows, including blank rows')
            if not any(v not in [None,''] for v in row):continue
            if book is not None and len(row)<len(headers):row=tuple(row)+(None,)*(len(headers)-len(row))
            if len(row)!=len(headers):raise AppError('TABLE_ROW','Every row must match the header columns')
            out.append(dict(zip(headers,row)))
            if len(out)>10000:raise AppError('ROW_LIMIT','Import at most 10,000 rows per file')
        if not out:raise AppError('EMPTY_TABLE','The table contains no data rows')
        return out
    except AppError:raise
    except (UnicodeError,csv.Error):raise AppError('CSV_FORMAT','Upload a UTF-8 CSV with a valid header and rows')
    except Exception:raise AppError('XLSX_FORMAT','Spreadsheet could not be read; use .xlsx without macros')
    finally:
        if book is not None:book.close()

def integer(value,default=None):
    if value is None or value=='':
        if default is None:raise ValueError('Missing integer value')
        value=default
    if isinstance(value,bool):raise ValueError('Boolean is not an integer amount')
    number=Decimal(str(value))
    if not number.is_finite() or number!=number.to_integral_value() or abs(number)>100000000:raise ValueError('Use whole numbers within 100,000,000')
    return int(number)


def parse_rule_text(text):
    # Parse complete labelled values, never a numeric prefix ("1" in "1e6")
    # or the first of conflicting facts. All results still require confirmation.
    fields={'kind':'policy'}
    if re.search(r'\bINR\b|₹',text,re.I):fields['currency']='INR'
    labels={
        'min_margin_bps':r'\b(?:minimum|min)[\s_-]*(?:contribution[\s_-]*)?margin\s*[:=]\s*',
        'max_discount_bps':r'\b(?:maximum|max)[\s_-]*discount\s*[:=]\s*',
        'budget_paise':r'\b(?:marketing[\s_-]*)?budget\s*[:=]\s*',
    }
    matches=sorted((m.start(),m.end(),key) for key,pattern in labels.items()
                   for m in re.finditer(pattern,text,re.I))
    observations={}
    for index,(_,end,key) in enumerate(matches):
        stop=matches[index+1][0] if index+1<len(matches) else len(text)
        token=re.split(r'[\r\n;]',text[end:stop],maxsplit=1)[0].strip()
        if token.endswith('.'):token=token[:-1].rstrip()
        value=None
        # Bound work and Decimal precision before processing attacker-controlled digits.
        if len(token)<=48:
            number=r'(\d[\d,]*(?:\.\d+)?|\.\d+)'
            pattern=(r'(?:INR|Rs\.?|₹)?\s*'+number+r'\s*(lakh|lakhs|lac|lacs|crore|crores)?'
                     if key=='budget_paise' else number+r'\s*%')
            parsed=re.fullmatch(pattern,token,re.I)
            if parsed:
                raw=parsed.group(1);whole=raw.split('.')[0]
                grouping=',' not in whole or re.fullmatch(r'(?:\d{1,3}(?:,\d{3})+|\d{1,2}(?:,\d{2})*,\d{3})',whole)
                if grouping and len(raw.replace(',',''))<=18:
                    scale=(parsed.group(2) or '').lower() if key=='budget_paise' else ''
                    multiplier=100000 if scale in ['lakh','lakhs','lac','lacs'] else (10000000 if scale in ['crore','crores'] else 1)
                    amount=Decimal(raw.replace(',',''))*multiplier*100
                    ceiling=100000000 if key=='budget_paise' else 10000
                    if 0<=amount<=ceiling and amount==amount.to_integral_value():value=int(amount)
        observations.setdefault(key,[]).append(value)
    for key,values in observations.items():
        fields[key]=values[0] if values[0] is not None and all(v==values[0] for v in values) else None
    if len(fields)==1:fields={'kind':'unresolved'}
    return fields

def ingest_document(store,tenant,name,content,mode='offline'):
    if mode not in ['offline','live']:raise AppError('MODE','Select offline or live')
    ext=Path(name).suffix.lower();text='';media=None
    if ext=='.pdf':
        try:
            reader=PdfReader(BytesIO(content))
            if len(reader.pages)>12:raise AppError('PAGE_LIMIT','Upload up to 12 pages at a time')
            text='\n'.join(f'Page {i+1}: {p.extract_text() or ""}' for i,p in enumerate(reader.pages))
            if len(text)>22000:raise AppError('TEXT_LIMIT','Extracted text exceeds 22,000 characters; split the source before review')
            media=('application/pdf',content)
        except AppError:raise
        except Exception:raise AppError('PDF_FORMAT','PDF could not be read')
    elif ext in ['.png','.jpg','.jpeg']:
        try:
            im=Image.open(BytesIO(content));im.verify();im=Image.open(BytesIO(content))
            if im.width*im.height>12_000_000:raise AppError('IMAGE_LIMIT','Image dimensions are too large')
            media=('image/png' if ext=='.png' else 'image/jpeg',content)
        except AppError:raise
        except Exception:raise AppError('IMAGE_FORMAT','Image could not be decoded')
    elif ext in ['.txt','.json','.md']:
        try:text=content.decode('utf8')
        except UnicodeError:raise AppError('TEXT_FORMAT','Text must use UTF-8 encoding')
        if len(text)>22000:raise AppError('TEXT_LIMIT','Text exceeds 22,000 characters; split the source before review')
    else:raise AppError('FILE_TYPE','Supported evidence: PDF, PNG, JPG, TXT, JSON, MD')
    provider=Provider(store,tenant)
    from .evidence_parser import extract
    if mode=='live':
        fields=provider.extract(text,media)
        provenance={key:{'value':value,'status':'model_extracted_unconfirmed','method':'provider','locations':[],
            'validation':'Model output only; exact source location and value require human review'} for key,value in fields.items() if key!='kind'}
    else:fields,provenance=extract(text,parse_rule_text)
    flagged=bool(re.search(r'ignore.*instruction|reveal.*secret|api.?key|override.*policy',text,re.I))
    record={'id':uid('src_'),'name':name,'sha256':digest(content),'kind':fields.get('kind','unresolved'),
       'fields':fields,'text':text,'source_size':len(content),'source_format':ext,'created_at':now(),'version':1,
       'mode':mode,'status':'needs_review','evidence_status':'extracted, unconfirmed', 'instruction_warning':flagged,
       'confidence_basis':'Schema and product/unit validation; human confirmation required before policy or pricing changes',
       'notes':['No automatic policy or price change. Review extracted values first.']}
    record['field_provenance']=provenance
    record['extraction_method']='provider' if mode=='live' else ('pdf_text' if ext=='.pdf' else 'explicit_text' if text else 'unsupported_raster')
    record['extraction_status']='review_required' if any(v is not None for k,v in fields.items() if k!='kind') else 'abstained'
    if mode=='offline' and (not text or not re.sub(r'Page \d+:|\s','',text)):
        record['notes'].append('No readable text layer. Local OCR is unavailable in this deployment; inspect the original and enter critical fields manually. No values were guessed.')
    unresolved=[key for key in ['min_margin_bps','max_discount_bps','budget_paise'] if key in fields and fields[key] is None]
    if unresolved:record['notes'].append('Unresolved, unsupported or conflicting labelled values require manual review: '+', '.join(unresolved))
    import base64
    record['source_base64']=base64.b64encode(content).decode()
    return store.put_new(tenant,'evidence',record['id'],record)

def import_table(store,tenant,name,content,kind):
    rows=table_rows(name,content)
    if len(rows)>10000:raise AppError('ROW_LIMIT','Import at most 10,000 rows per file')
    state=store.get(tenant,'state','retailer');products={p['sku']:p for p in state['products']};updates=[];errors=[]
    seen=set()
    for no,r in enumerate(rows,2):
        try:
            if r.get('currency','INR')!='INR':raise ValueError('Only INR operational data is supported')
            if kind=='inventory':
                sku=str(r.get('sku','')).strip();city=str(r.get('city','')).strip()
                if sku not in products or city not in CITIES:raise ValueError('Unknown SKU or city')
                basis=str(r.get('quantity_basis','physical'));quantity=integer(r['quantity']);reserved=integer(r.get('reserved'),0);unavailable=integer(r.get('unavailable'),0);safety=integer(r.get('safety_stock'),2)
                if min(quantity,reserved,unavailable,safety)<0:raise ValueError('Quantities must be nonnegative')
                if basis not in ['physical','available_to_promise']:raise ValueError('Use physical or available_to_promise quantity basis')
                if basis=='available_to_promise':quantity+=reserved+unavailable+safety
                expiry=str(r['expiry'])[:10];datetime.fromisoformat(expiry)
                observed=str(r.get('observed_at') or now());dt=datetime.fromisoformat(observed)
                if dt.tzinfo is None or dt>datetime.now(timezone.utc)+timedelta(seconds=60):raise ValueError('Timestamp needs a UTC offset and cannot be in the future')
                update={'sku':sku,'city':city,'quantity':quantity,'reserved':reserved,'unavailable':unavailable,'safety_stock':safety,'expiry':expiry,'observed_at':observed}
            elif kind=='competitor':
                sku=str(r['sku']);city=str(r['city'])
                if sku not in products or city not in CITIES:raise ValueError('Unknown SKU or city')
                if r.get('currency','INR')!='INR':raise ValueError('Only INR observations are supported')
                price=integer(r['price_paise']);pack=integer(r['pack_quantity']);unit=str(r['pack_unit']);observed=str(r['observed_at']);dt=datetime.fromisoformat(observed)
                if dt.tzinfo is None or dt>datetime.now(timezone.utc)+timedelta(seconds=60) or price<=0 or pack<=0:raise ValueError('Invalid price, pack, or timestamp')
                update={'sku':sku,'city':city,'price_paise':price,'pack_quantity':pack,'pack_unit':unit,'observed_at':observed,'conditions':str(r.get('conditions',''))}
            elif kind=='products':
                sku=str(r['sku']);p=products.get(sku)
                if not p:raise ValueError('Import existing SKU costs; catalog extension requires mapping')
                cost=integer(r['cost_paise']);price=integer(r['price_paise'])
                if cost<0 or price<=0:raise ValueError('Invalid cost or price')
                update={'sku':sku,'cost_paise':cost,'price_paise':price}
            elif kind=='transactions':
                sku=str(r['sku']);city=str(r['city']);day=datetime.fromisoformat(str(r['date'])[:10]).date()
                if sku not in products or city not in CITIES:raise ValueError('Unknown SKU or city')
                anchor=datetime.fromisoformat(state.get('history_start_date',str(business_date(datetime.fromisoformat(state['created_at']))-timedelta(days=180)))).date()
                index=(day-anchor).days
                if not 0<=index<180 or day>=business_date():raise ValueError('Transaction date must be a complete IST day in the snapshot history window')
                quantity=integer(r['quantity']);paid=integer(r['paid_price_paise']);flag=str(r.get('stockout','false')).lower()
                if quantity<0 or paid<=0 or flag not in ['true','false','1','0']:raise ValueError('Use nonnegative sold quantity, positive per-item net price and true/false stockout')
                items=[x.strip()for x in str(r.get('basket_items','')).split('|')if x.strip()]
                if set(items)-set(products):raise ValueError('Basket includes an unmapped SKU')
                update={'sku':sku,'city':city,'date':day.isoformat(),'day_index':index,'quantity':quantity,'paid_price_paise':paid,'stockout':flag in ['true','1'],'weekday':day.weekday(),'basket_items':items}
            else:raise ValueError('Import kind must be inventory, competitor, products or transactions')
            identity=(sku,city,day.isoformat()) if kind=='transactions' else ((sku,city) if kind in ['inventory','competitor'] else (sku,))
            if identity in seen:raise ValueError('Duplicate product/location/day row; aggregate or reconcile before importing')
            seen.add(identity)
            updates.append(update)
        except (ValueError,KeyError,TypeError,ArithmeticError) as e:errors.append({'row':no,'message':str(e)})
    record={'id':'imp_'+digest([kind,digest(content)])[:28],'name':name,'kind':kind,'rows':len(rows),'valid_rows':len(updates),'errors':errors,'status':'rejected' if errors else 'applied','created_at':now(),'sha256':digest(content),'schema_version':1}
    if not errors:
        def apply(s,conn):
            existing=store.read_in(conn,tenant,'import',record['id'])
            if existing:return existing
            if kind=='products' and any(p['status']=='active' for p in store.list_in(conn,tenant,'plan')):raise AppError('ACTIVE_COMMITMENT','Pause active offers before changing their financial inputs',409)
            for u in updates:
                if kind=='inventory':
                    b=next(b for b in s['inventory'] if b['sku']==u['sku'] and b['city']==u['city'])
                    check_inventory_observation(store,tenant,conn,b,datetime.fromisoformat(u['observed_at']))
                    if b['reserved'] and u['expiry']!=b['expiry']:raise AppError('ACTIVE_RESERVATION','Pause active offers before changing committed batch expiry',409)
                    if b['reserved'] and (u['quantity']-u['unavailable']-u['safety_stock']<b['reserved'] or u['reserved']!=b['reserved']):raise AppError('ACTIVE_RESERVATION','Pause active offers before replacing committed stock reservations',409)
                    b.update({k:v for k,v in u.items() if k not in ['quantity']});b['physical_on_hand']=u['quantity'];b['version']+=1
                elif kind=='competitor':
                    s['competitors'].append({'id':uid('cmp_'),**u,'source':record['id'],'evidence_status':'observed','version':1})
                elif kind=='transactions':
                    s['history'][u['sku']+'|'+u['city']][u['day_index']]=[u['quantity'],u['paid_price_paise'],u['stockout'],u['weekday']]
                    if len(u['basket_items'])>=2:s['baskets'].append({'id':record['id']+'-'+str(len(s['baskets'])),'items':u['basket_items'],'source':record['id']})
                    s['history_source']=record['id'];s['data_label']='Mixed imported and demo store data'
                else:
                    p=next(p for p in s['products'] if p['sku']==u['sku']);p.update(u);p['version']+=1;p['source']=record['id'];p['evidence_status']='observed'
            s['data_label']='Mixed imported and demo store data'
            s['stock_version']+=1;s['snapshot_id']=digest([s['snapshot_id'],record['id']])[:16]
            store.write_in(conn,tenant,'import',record['id'],record)
            return record
        return store.atomic(tenant,apply)
    return store.put_new(tenant,'import',record['id'],record)

def correct_evidence(store,tenant,id,fields):
    allowed={'kind','sku','city','name','currency','price_paise','pack_quantity','pack_unit','min_margin_bps','max_discount_bps','budget_paise','conditions','observed_at','source_quote'}
    if not isinstance(fields,dict):raise AppError('CORRECTION_FORMAT','Corrections must be an object')
    if set(fields)-allowed:raise AppError('FIELD_NOT_ALLOWED','Correction contains unsupported fields')
    previous=store.get(tenant,'evidence',id);new={**previous,'id':uid('src_'),'parent_id':id,'version':previous['version']+1,'fields':{**previous['fields'],**fields},'status':'confirmed','evidence_status':'human confirmed','created_at':now()}
    new['field_provenance']={k:{**previous.get('field_provenance',{}).get(k,{}),'value':v,'status':'human_confirmed',
        'confirmation':{'previous_id':id,'changed':v!=previous['fields'].get(k),'at':new['created_at']}} for k,v in new['fields'].items() if k!='kind'}
    new['extraction_status']='human_confirmed'
    f=new['fields'];state=store.get(tenant,'state','retailer')
    if f.get('currency')!='INR':raise AppError('CURRENCY','Confirm INR evidence; unsupported currency cannot be applied')
    try:
        for key in ['price_paise','pack_quantity','min_margin_bps','max_discount_bps','budget_paise']:
            if f.get(key) is not None:f[key]=integer(f[key])
        for key in ['kind','sku','city','currency','pack_unit','observed_at','conditions']:
            if f.get(key) is not None and not isinstance(f[key],str):raise ValueError('Text fields must be strings')
    except (ValueError,ArithmeticError,TypeError):raise AppError('CORRECTION_FORMAT','Use valid whole-number amounts and text fields')
    new['kind']=f.get('kind','unresolved')
    def apply(s,conn):
        if f.get('kind')=='policy':
            if not all(f.get(k) is not None for k in ['min_margin_bps','max_discount_bps','budget_paise']):raise AppError('INCOMPLETE_POLICY','Confirm margin, discount and budget before applying')
            if not 0<=int(f['min_margin_bps'])<=9000 or not 0<=int(f['max_discount_bps'])<=5000 or int(f['budget_paise'])<0:raise AppError('INVALID_POLICY','Policy values are outside supported bounds')
            if any(p['status']=='active' for p in store.list_in(conn,tenant,'plan')):raise AppError('ACTIVE_COMMITMENT','Pause active offers before replacing the policy',409)
            if f['budget_paise']<s['spent_paise']:raise AppError('BUDGET_BELOW_SPEND','The policy budget cannot erase already incurred spend',409)
            s['policy'].update({'min_margin_bps':int(f['min_margin_bps']),'max_discount_bps':int(f['max_discount_bps']),'marketing_budget_paise':int(f['budget_paise']),'version':s['policy']['version']+1,'source':new['id'],'status':'human confirmed','observed_at':now()})
        elif f.get('kind')=='competitor':
            p=next((p for p in s['products'] if p['sku']==f.get('sku')),None)
            if not p or f.get('city')not in CITIES:raise AppError('PRODUCT_MATCH','Confirm an existing SKU and supported city')
            if any(f.get(k)is None for k in ['currency','price_paise','pack_quantity','pack_unit','observed_at']):raise AppError('INCOMPLETE_OFFER','Currency, price, pack, unit and timestamp must be confirmed')
            if p['pack_quantity']!=int(f['pack_quantity']) or p['pack_unit']!=f['pack_unit']:raise AppError('PACK_MISMATCH','Pack size does not match the product')
            if int(f['price_paise'])<=0:raise AppError('INVALID_PRICE','Price must be positive')
            try:dt=datetime.fromisoformat(f['observed_at'])
            except ValueError:raise AppError('TIMESTAMP','Use a valid ISO timestamp')
            if dt.tzinfo is None or dt>datetime.now(timezone.utc)+timedelta(seconds=60):raise AppError('TIMESTAMP','Timestamp requires timezone and cannot be in the future')
            s['competitors'].append({'id':new['id'],'sku':p['sku'],'city':f['city'],'price_paise':int(f['price_paise']),'pack_quantity':p['pack_quantity'],'pack_unit':p['pack_unit'],'observed_at':f['observed_at'],'source':new['id'],'conditions':f.get('conditions',''),'evidence_status':'human confirmed','version':new['version']})
        else:raise AppError('EVIDENCE_KIND','Choose policy or competitor before applying')
        s['snapshot_id']=digest([s['snapshot_id'],new['id']])[:16];s['stock_version']+=1
        store.write_in(conn,tenant,'evidence',new['id'],new)
    store.atomic(tenant,apply);return new
